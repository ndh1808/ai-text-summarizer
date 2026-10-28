import sys
import unittest
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from document_reader import DocumentError, extract_document, MAX_UPLOAD_BYTES, MAX_TEXT_CHARS, MAX_PDF_PAGES


class DocumentReaderTests(unittest.TestCase):
    def test_utf8_text(self):
        result = extract_document("sample.txt", "Xin chào Việt Nam".encode("utf-8"))
        self.assertEqual(result.text, "Xin chào Việt Nam")
        self.assertFalse(result.truncated)

    def test_invalid_extension(self):
        with self.assertRaises(DocumentError):
            extract_document("bad.exe", b"Hello")

    def test_empty_file(self):
        with self.assertRaises(DocumentError):
            extract_document("empty.pdf", b"")

    def test_file_too_large(self):
        with self.assertRaises(DocumentError):
            extract_document("large.txt", b"x" * (MAX_UPLOAD_BYTES + 1))

    def test_truncation_notice(self):
        doc = extract_document("long.md", b"x" * (MAX_TEXT_CHARS + 123))
        self.assertTrue(doc.truncated)
        self.assertEqual(len(doc.text), MAX_TEXT_CHARS)

    def test_limits_raised(self):
        self.assertEqual(MAX_UPLOAD_BYTES, 50 * 1024 * 1024)
        self.assertEqual(MAX_PDF_PAGES, 300)
        self.assertEqual(MAX_TEXT_CHARS, 200_000)

    def test_docx_paragraphs_and_table(self):
        from docx import Document
        document = Document()
        document.add_paragraph("Mở đầu tài liệu")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Tiêu đề"
        table.cell(0, 1).text = "Giá trị"
        document.add_paragraph("Kết thúc tài liệu")
        buffer = BytesIO()
        document.save(buffer)
        extracted = extract_document("bai_hoc.DOCX", buffer.getvalue()).text
        self.assertIn("Mở đầu tài liệu", extracted)
        self.assertIn("Tiêu đề | Giá trị", extracted)
        self.assertIn("Kết thúc tài liệu", extracted)
        self.assertLess(extracted.index("Mở đầu"), extracted.index("Tiêu đề"))
        self.assertLess(extracted.index("Tiêu đề"), extracted.index("Kết thúc"))

    def test_blank_pdf_scan_message(self):
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        buffer = BytesIO()
        writer.write(buffer)
        with self.assertRaisesRegex(DocumentError, "OCR"):
            extract_document("scan.pdf", buffer.getvalue())

    def test_pdf_text(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        writer = PdfWriter()
        page = writer.add_blank_page(width=400, height=300)
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        font_ref = writer._add_object(font)
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})
        })
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 20 200 Td (Hello PDF document.) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        buffer = BytesIO()
        writer.write(buffer)
        output = extract_document("example.pdf", buffer.getvalue())
        self.assertIn("Hello PDF document", output.text)

    def test_password_or_corrupt_pdf(self):
        with self.assertRaises(DocumentError):
            extract_document("broken.pdf", b"%PDF-corrupt")


if __name__ == "__main__":
    unittest.main()
